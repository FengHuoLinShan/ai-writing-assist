"""表格解析契约（计划 §3.2）— 有界解析 xlsx / csv，同步函数。

调用方（L4 service）通过 ``asyncio.to_thread`` 调用。所有限额按 constants.py
执行，超限一律明确拒绝（抛作者可读 ValueError），不静默截断。
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from modules.imports.parsers import audit_zip_container, detect_encoding
from modules.imports.spreadsheet_migration.constants import (
    MAX_CELL_CHARS,
    MAX_COLUMNS,
    MAX_ROWS_PER_SHEET,
    MAX_SHEETS_PER_FILE,
    XLSX_MAX_MEMBERS,
    XLSX_MAX_UNCOMPRESSED,
)


@dataclass(frozen=True)
class ParsedSheet:
    """一张解析后的工作表；rows 全部为字符串，合并单元格已填充。"""

    sheet_key: str  # f"f{file_idx}s{sheet_idx}"，会话内唯一
    file_key: str  # f"f{file_idx}"
    name: str  # 表名；csv 用去扩展名的文件名
    hidden: bool
    rows: list[list[str]]
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedUpload:
    """一次上传解析出的全部工作表。"""

    file_key: str
    file_name: str
    file_type: str
    size: int
    sha256: str
    sheets: list[ParsedSheet] = field(default_factory=list)


__all__ = [
    "MAX_CELL_CHARS",
    "MAX_COLUMNS",
    "MAX_ROWS_PER_SHEET",
    "MAX_SHEETS_PER_FILE",
    "ParsedSheet",
    "ParsedUpload",
    "parse_xlsx",
    "parse_csv",
]


_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_XLSX_MEMBER_BYTES = XLSX_MAX_UNCOMPRESSED  # 单成员上限与解压总量一致
_XML_MEMBER_SCAN_BYTES = 1024 * 1024  # DOCTYPE/ENTITY 只需扫描文件头
_XML_MEMBER_SUFFIXES = (".xml", ".rels")
_NOTION_URL_SUFFIX = re.compile(
    r"\s*\(\s*(?:https?://)?(?:www\.)?notion\.so[^)]*\)\s*$",
    re.IGNORECASE,
)
_NOTION_HEX_SUFFIX = re.compile(r"\s+[0-9a-fA-F]{32}\s*$")
_CSV_DELIMITER_CANDIDATES = ",\t;"
_CSV_SNIFF_SAMPLE_BYTES = 64 * 1024


class _CellTooLongError(ValueError):
    """内部信号：单元格超过 MAX_CELL_CHARS，由调用处换成作者可读文案。"""


def _reject(message: str) -> ValueError:
    return ValueError(message)


def _strip_notion_suffixes(text: str) -> str:
    """去掉 Notion 导出残留的链接后缀和 32 位 hex 后缀。"""
    stripped = _NOTION_URL_SUFFIX.sub("", text).strip()
    stripped = _NOTION_HEX_SUFFIX.sub("", stripped).strip()
    return stripped


def _strip_trailing_empty(rows: list[list[str]]) -> list[list[str]]:
    """去掉尾部全空行和尾部全空列。"""
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    while width > 0 and all(not row[width - 1].strip() for row in rows):
        width -= 1
    return [row[:width] for row in rows]


def _enforce_sheet_limits(rows: list[list[str]], sheet_label: str) -> list[list[str]]:
    rows = _strip_trailing_empty(rows)
    if rows and len(rows) > MAX_ROWS_PER_SHEET:
        raise _reject(
            f"「{sheet_label}」数据行数超过上限（最多 {MAX_ROWS_PER_SHEET} 行），"
            "请拆分后再上传"
        )
    width = max((len(row) for row in rows), default=0)
    if width > MAX_COLUMNS:
        raise _reject(
            f"「{sheet_label}」列数超过上限（最多 {MAX_COLUMNS} 列），请拆分后再上传"
        )
    return rows


def _check_cell_length(text: str) -> str:
    if len(text) > MAX_CELL_CHARS:
        raise _CellTooLongError(str(MAX_CELL_CHARS))
    return text


# ============================================================
# xlsx
# ============================================================


def _xlsx_cell_text(value: object) -> str:
    """把 openpyxl 单元格值转成作者可读字符串。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, dt.datetime):
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e16:
            return str(int(value))
        return str(value)
    return str(value)


def _xlsx_formula_marker(value: object) -> bool:
    """判断 data_only=False 工作簿里的单元格是否为公式。"""
    if isinstance(value, str):
        return value.startswith("=")
    formula_text = getattr(value, "text", None)
    return isinstance(formula_text, str) and formula_text.startswith("=")


def _audit_xlsx_container(data: bytes) -> None:
    """xlsx 专属容器审计：签名、必需部件、宏与不安全 XML。"""
    if data[:8] == _OLE_SIGNATURE:
        raise _reject(
            "这是旧版 Excel（.xls）或加密表格，请在 Excel/WPS 中另存为 .xlsx 后重新上传"
        )
    audit_zip_container(
        data,
        max_members=XLSX_MAX_MEMBERS,
        max_member_bytes=_XLSX_MEMBER_BYTES,
        max_total_uncompressed=XLSX_MAX_UNCOMPRESSED,
    )
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = {info.filename.casefold() for info in archive.infolist()}
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise _reject(
                    "文件内容与 .xlsx 格式不符，请在 Excel/WPS 中另存为 .xlsx 后重新上传"
                )
            for name in names:
                if name.replace("\\", "/").endswith("vbaproject.bin"):
                    raise _reject(
                        "表格包含宏，为安全起见不支持，请另存为普通 .xlsx 文件后重新上传"
                    )
            for info in archive.infolist():
                lowered_name = info.filename.casefold().replace("\\", "/")
                content_types_member = lowered_name == "[content_types].xml"
                is_xml_member = lowered_name.endswith(_XML_MEMBER_SUFFIXES)
                if not (content_types_member or is_xml_member):
                    continue
                scan_cap = min(info.file_size, _XML_MEMBER_SCAN_BYTES)
                with archive.open(info, "r") as stream:
                    head = stream.read(scan_cap)
                lowered = head.lower()
                if b"<!doctype" in lowered or b"<!entity" in lowered:
                    raise _reject("文件内容与扩展名不匹配")
                if content_types_member and b"macroenabled" in lowered:
                    raise _reject(
                        "表格包含宏，为安全起见不支持，请另存为普通 .xlsx 文件后重新上传"
                    )
    except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
        raise _reject(
            "文件内容与 .xlsx 格式不符，请在 Excel/WPS 中另存为 .xlsx 后重新上传"
        ) from exc


def _xlsx_has_formulas(data: bytes) -> bool:
    """快速探测工作表部件是否含公式，决定是否需要第二本公式工作簿。"""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for info in archive.infolist():
                if not info.filename.casefold().endswith(".xml"):
                    continue
                with archive.open(info, "r") as stream:
                    head = stream.read(_XML_MEMBER_SCAN_BYTES)
                if b"<f>" in head or b"<f " in head:
                    return True
    except (OSError, RuntimeError, zipfile.BadZipFile):
        return False
    return False


def parse_xlsx(data: bytes, *, file_key: str, file_name: str) -> ParsedUpload:
    """有界解析 xlsx 字节流；违规抛 ValueError（作者可读）。"""
    if not data:
        raise _reject("文件没有内容，无法识别表格")
    _audit_xlsx_container(data)

    needs_formula_book = _xlsx_has_formulas(data)
    try:
        cached_book = load_workbook(
            io.BytesIO(data),
            read_only=False,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        formula_book = (
            load_workbook(
                io.BytesIO(data),
                read_only=False,
                data_only=False,
                keep_vba=False,
                keep_links=False,
            )
            if needs_formula_book
            else None
        )
    except (
        InvalidFileException,
        OSError,
        RuntimeError,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        TypeError,
        AttributeError,
        IndexError,
        SyntaxError,
    ) as exc:
        raise _reject(
            "表格文件无法解析，请确认它是 Excel/WPS 保存的正常 .xlsx 文件"
        ) from exc

    if len(cached_book.worksheets) > MAX_SHEETS_PER_FILE:
        raise _reject(
            f"工作表数量超过上限（单个文件最多 {MAX_SHEETS_PER_FILE} 张），请拆分后再上传"
        )

    sheets: list[ParsedSheet] = []
    try:
        for sheet_idx, worksheet in enumerate(cached_book.worksheets):
            sheets.append(
                _build_parsed_sheet(
                    worksheet,
                    formula_book[worksheet.title] if formula_book else None,
                    sheet_idx=sheet_idx,
                    file_key=file_key,
                )
            )
    except _CellTooLongError as exc:
        raise _reject(
            f"某个单元格内容超过 {MAX_CELL_CHARS} 字上限，请缩短或拆分该单元格"
        ) from exc

    return ParsedUpload(
        file_key=file_key,
        file_name=file_name,
        file_type="xlsx",
        size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        sheets=sheets,
    )


def _build_parsed_sheet(
    worksheet,
    formula_worksheet,
    *,
    sheet_idx: int,
    file_key: str,
) -> ParsedSheet:
    warnings: list[str] = []
    hidden = str(worksheet.sheet_state or "visible") != "visible"
    sheet_name = str(worksheet.title or f"工作表{sheet_idx + 1}")

    merged_values: dict[tuple[int, int], object] = {}
    for merged_range in worksheet.merged_cells.ranges:
        top_left = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        if top_left is None:
            continue
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                if (row, column) != (merged_range.min_row, merged_range.min_col):
                    merged_values[(row, column)] = top_left

    image_count = len(getattr(worksheet, "_images", None) or [])
    comment_count = 0
    formula_without_cache = False

    cached_rows = worksheet.iter_rows()
    formula_rows = (
        iter(formula_worksheet.iter_rows()) if formula_worksheet is not None else None
    )

    rows: list[list[str]] = []
    for row_cells in cached_rows:
        formula_cells = next(formula_rows, None) if formula_rows is not None else None
        row_text: list[str] = []
        for column_idx, cell in enumerate(row_cells, start=1):
            if cell.comment is not None:
                comment_count += 1
            value = cell.value
            if value is None:
                value = merged_values.get((cell.row, cell.column))
            if value is None and formula_cells is not None:
                sibling = (
                    formula_cells[column_idx - 1]
                    if column_idx - 1 < len(formula_cells)
                    else None
                )
                if sibling is not None and _xlsx_formula_marker(sibling.value):
                    formula_without_cache = True
                    value = ""
            row_text.append(_check_cell_length(_xlsx_cell_text(value)))
        rows.append(row_text)

    if formula_without_cache:
        warnings.append("部分公式没有保存计算结果，对应单元格已按空值处理")
    if image_count or comment_count:
        parts = []
        if image_count:
            parts.append(f"{image_count} 张图片")
        if comment_count:
            parts.append(f"{comment_count} 条批注")
        warnings.append("已忽略：" + "、".join(parts))

    return ParsedSheet(
        sheet_key=f"{file_key}s{sheet_idx}",
        file_key=file_key,
        name=sheet_name,
        hidden=hidden,
        rows=_enforce_sheet_limits(rows, sheet_name),
        warnings=warnings,
    )


# ============================================================
# csv
# ============================================================


def _decode_csv_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise _reject(
                "表格编码无法识别，请使用 UTF-8 或 GBK 编码保存后重新上传"
            ) from exc
    encoding = detect_encoding(data).lower()
    if encoding in {"gb2312", "gbk"}:
        encoding = "gb18030"
    try:
        return data.decode(encoding, errors="strict")
    except (UnicodeDecodeError, LookupError):
        for fallback in ("utf-8", "gb18030"):
            try:
                return data.decode(fallback)
            except UnicodeDecodeError:
                continue
        raise _reject(
            "表格编码无法识别，请使用 UTF-8 或 GBK 编码保存后重新上传"
        ) from None


def _sniff_csv_delimiter(text: str) -> str:
    sample = text[:_CSV_SNIFF_SAMPLE_BYTES]
    if not sample.strip():
        return ","
    try:
        return csv.Sniffer().sniff(sample, delimiters=_CSV_DELIMITER_CANDIDATES).delimiter
    except csv.Error:
        return ","


def parse_csv(data: bytes, *, file_key: str, file_name: str) -> ParsedUpload:
    """有界解析 csv 字节流；违规抛 ValueError（作者可读）。"""
    text = _decode_csv_text(data)
    delimiter = _sniff_csv_delimiter(text)
    csv.field_size_limit(MAX_CELL_CHARS + 1)

    rows: list[list[str]] = []
    try:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        for raw_row in reader:
            cells = [_strip_notion_suffixes(str(cell)) for cell in raw_row]
            try:
                cells = [_check_cell_length(cell) for cell in cells]
            except _CellTooLongError as exc:
                raise _reject(
                    f"某个单元格内容超过 {MAX_CELL_CHARS} 字上限，请缩短或拆分该单元格"
                ) from exc
            rows.append(cells)
    except csv.Error as exc:
        raise _reject(
            f"某个单元格内容超过 {MAX_CELL_CHARS} 字上限，请缩短或拆分该单元格"
        ) from exc

    base_name = file_name.replace("\\", "/").rsplit("/", 1)[-1]
    sheet_name = _strip_notion_suffixes(base_name.rsplit(".", 1)[0]).strip() or "表格"

    rows = _enforce_sheet_limits(rows, sheet_name)

    return ParsedUpload(
        file_key=file_key,
        file_name=file_name,
        file_type="csv",
        size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        sheets=[
            ParsedSheet(
                sheet_key=f"{file_key}s0",
                file_key=file_key,
                name=sheet_name,
                hidden=False,
                rows=rows,
                warnings=[],
            )
        ],
    )
