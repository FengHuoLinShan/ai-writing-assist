"""表格迁移测试夹具 — 用 openpyxl 在测试内生成 xlsx/csv 字节，不落盘。"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.comments import Comment


def build_xlsx(sheets: list[dict[str, Any]]) -> bytes:
    """生成 xlsx 字节。

    sheets 里每项支持：
      name: 表名（默认 Sheet{n}）
      rows: list[list[value]]，按行追加（跳过 None 标题占位由调用方控制）
      cell: (row, column, value) 精确写入，1 起
      merges: ["A1:B2", ...] 合并单元格
      hidden: True 标记隐藏表
      comment: (row, column, text) 给单元格加批注
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, spec in enumerate(sheets):
        worksheet = workbook.create_sheet(title=spec.get("name") or f"Sheet{index + 1}")
        if spec.get("hidden"):
            worksheet.sheet_state = "hidden"
        for row in spec.get("rows", []):
            worksheet.append(row)
        for row, column, value in spec.get("cell", []):
            worksheet.cell(row=row, column=column, value=value)
        for merge in spec.get("merges", []):
            worksheet.merge_cells(merge)
        for row, column, text in spec.get("comment", []):
            target: Cell = worksheet.cell(row=row, column=column)
            target.comment = Comment(text, "测试")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def inject_formula_cache(data: bytes, cached_value: str = "42") -> bytes:
    """给 openpyxl 生成的公式单元格补写缓存值 <v>。

    openpyxl 只写 <f> 和空 <v></v>；真实 Excel 会带缓存值。这里把空 <v>
    替换成具体值（没有 <v> 时补写），模拟“上次由 Excel 计算保存”。
    """
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            content = source.read(info)
            if info.filename.startswith("xl/worksheets/") and info.filename.endswith(
                ".xml"
            ):
                text = content.decode("utf-8")
                text = re.sub(
                    r"(<f[^>]*>[^<]*</f>)(?:<v>\s*</v>)?",
                    rf"\1<v>{cached_value}</v>",
                    text,
                )
                content = text.encode("utf-8")
            target.writestr(info, content)
    return output.getvalue()


def rebuild_zip_with_extra_members(
    data: bytes,
    extra_members: list[tuple[str, bytes]],
) -> bytes:
    """在既有 xlsx zip 后追加成员（用于构造路径穿越、多余成员等恶意样本）。"""
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            target.writestr(info, source.read(info))
        for name, content in extra_members:
            target.writestr(name, content)
    return output.getvalue()


def replace_member(
    data: bytes,
    member_name: str,
    content: bytes,
) -> bytes:
    """替换 xlsx zip 中指定成员的内容。"""
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            if info.filename == member_name:
                target.writestr(info, content)
            else:
                target.writestr(info, source.read(info))
    return output.getvalue()


def build_zip_bomb_xlsx(*, declared_total_uncompressed: int) -> bytes:
    """构造解压总量超限的“xlsx”：合法部件 + 一个高压缩比大成员。"""
    base = build_xlsx([{"name": "人物", "rows": [["姓名"], ["张三"]]}])
    bomb_payload = b"A" * (declared_total_uncompressed // 2)
    return rebuild_zip_with_extra_members(base, [("xl/bigBlob.bin", bomb_payload)])


def build_member_bomb_xlsx(member_count: int) -> bytes:
    """构造成员数超限的“xlsx”。"""
    base = build_xlsx([{"name": "人物", "rows": [["姓名"], ["张三"]]}])
    extra = [(f"xl/extra{i}.bin", b"x") for i in range(member_count)]
    return rebuild_zip_with_extra_members(base, extra)


OLE_BYTES = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512

NOTION_HEX_SUFFIX = "48e2f9c8d1a2b3c4e5f6a7b8c9d0e1f2"
